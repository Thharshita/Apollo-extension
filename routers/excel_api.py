import pandas as pd
import io
import asyncio
from fastapi import APIRouter
from fastapi import UploadFile, File
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side
from fastapi.responses import StreamingResponse, JSONResponse
from logger import logger
from constant import PHONE_COL, EMAIL_COL, COMPANY_WEBSITE_COL, COMPANY_NAME_COL
from scrapers.google_scraper import google_search
from scrapers.web_scraper_bs import fetch_links, extract_contents
from utils import process_tat
import time

router = APIRouter(tags=["Excel Processing"])
EXCEL_BATCH_SIZE = 5


@router.post("/process-excel/")
async def process_excel(file: UploadFile):
    try:
        start_time = time.time()
        contents = await file.read()
        if file.filename.endswith('.csv'):
            df = pd.read_csv(io.BytesIO(contents))
        else:
            df = pd.read_excel(io.BytesIO(contents))
        logger.info(f"File loaded: {len(df)} rows, columns: {list(df.columns)}")

        _insert_col_after(df, EMAIL_COL, "Scraped Email")
        # _insert_col_after(df, "Google Email", "Scraped Email")
        _insert_col_after(df, PHONE_COL, "Scraped Phone")
        # _insert_col_after(df, "Google Phone", "Scraped Phone")
        _insert_col_after(df, COMPANY_WEBSITE_COL, "Scraped Website")


        rows_to_process = [
            (idx, row.to_dict())
            for idx, row in df.iterrows()
            if _is_missing(row.get(EMAIL_COL)) or _is_missing(row.get(PHONE_COL))
        ]

        metrics=[]

        for batch_start in range(0, len(rows_to_process), EXCEL_BATCH_SIZE):
            batch = rows_to_process[batch_start:batch_start + EXCEL_BATCH_SIZE]
            results = await asyncio.gather(
                *(_process_row(idx, row) for idx, row in batch)
            )
            metrics.append(results)

            for idx, scraped, google_result, alternate_website_scraped in results:
                if scraped.get("email"):
                    df.at[idx, "Scraped Email"] = scraped["email"]
                if scraped.get("number"):
                    df.at[idx, "Scraped Phone"] = scraped["number"]
                if scraped.get("scraped_from_link"):
                    df.at[idx, "Scraped Website"] = scraped["scraped_from_link"]
                if google_result.get("email") and not df.at[idx, "Scraped Email"]:
                    df.at[idx, "Scraped Email"] = google_result["email"]
                if google_result.get("number") and not df.at[idx, "Scraped Phone"]:
                    df.at[idx, "Scraped Phone"] = google_result["number"]
                if google_result.get("website") and not df.at[idx, "Scraped Website"]:
                    df.at[idx, "Scraped Website"] = google_result["website"]
                if alternate_website_scraped.get("email") and not df.at[idx, "Scraped Email"]:
                    df.at[idx, "Scraped Email"] = alternate_website_scraped["email"]
                if alternate_website_scraped.get("number") and not df.at[idx, "Scraped Phone"]:
                    df.at[idx, "Scraped Phone"] = alternate_website_scraped["number"]
                if alternate_website_scraped.get("scraped_from_link") and not df.at[idx, "Scraped Website"]:
                    df.at[idx, "Scraped Website"] = alternate_website_scraped["scraped_from_link"]

        NEW_COLS = {"Scraped Email", "Scraped Phone", "Scraped Website"}
        output = io.BytesIO()
        # if file.filename.endswith('.csv'):
        #     df.to_csv(output, index=False)
        #     media_type = "text/csv"
        # else:
        #providing excel as ouput because csv cant hold colors.
        logger.info(f"metrics: {metrics}")
        df.to_excel(output, index=False)

        _highlight_columns(output, df.columns.tolist(), NEW_COLS)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        output.seek(0)
        end_time=time.time()
        tat=process_tat(start_time,end_time)
        logger.info(f"Total time taken: {tat/60} minutes")
        return StreamingResponse(
            output,
            media_type=media_type,
            headers={"Content-Disposition": f"attachment; filename=updated_{file.filename}"}
        )

    except Exception as e:
        logger.error(f"process_excel failed: {e}")
        return JSONResponse(status_code=500, content={"status": 0, "message": str(e)})


def _is_missing(val) -> bool:
    if val is None:
        return True
    if pd.isna(val):
        return True
    if isinstance(val, str) and val.strip().strip("'").lower() in ("", "nan", "none", "n/a", "na"):
        return True
    return False


def _insert_col_after(df: pd.DataFrame, after_col: str, new_col: str):
    """Insert new_col with None values immediately after after_col."""
    if after_col in df.columns:
        pos = df.columns.get_loc(after_col) + 1
    else:
        pos = len(df.columns)
    df.insert(pos, new_col, None)


async def _process_row(idx, row: dict):
    company_name = row.get(COMPANY_NAME_COL, "")
    email_missing = _is_missing(row.get(EMAIL_COL))
    phone_missing = _is_missing(row.get(PHONE_COL))
    website = row.get(COMPANY_WEBSITE_COL)
    scraped = {}
    google_result = {}
    alternate_website_scraped = {}
    scraped_from_google_website = {}

    logger.info(f"Processing row {idx}: {company_name}")

    if not _is_missing(website):
        logger.info(f"metrics: Scraping website for {company_name} at {website}")
        scraped = await asyncio.to_thread(
            asyncio.run,
            _scrape_website(str(website)),
        )
        logger.info(f"metrics: Scraped result for {company_name}: {scraped}")
        email_missing = email_missing and not scraped.get("email") 
        phone_missing = phone_missing and not scraped.get("number")

    if email_missing or phone_missing:
        logger.info(f"metrics: Google searching for {company_name}")
        google_result = await asyncio.to_thread(
            google_search,
            str(company_name),
            known_website=None if _is_missing(website) else str(website),
        )
        if google_result:
            email_missing = email_missing and not google_result.get("email")
            phone_missing = phone_missing and not google_result.get("number")
            website_from_google = google_result.get("website")

            if (email_missing or phone_missing) and website_from_google:
                logger.info(f"metrics: Scraping website from Google result for {company_name} at {website_from_google}")
                scraped_from_google_website = await asyncio.to_thread(
                    asyncio.run,
                    _scrape_website(str(website_from_google)),
                )
                logger.info(f"metrics: Scraped result from Google website for {company_name}: {scraped_from_google_website}")
                if scraped_from_google_website.get("email"):
                    alternate_website_scraped["email"] = scraped_from_google_website["email"]
                if scraped_from_google_website.get("number"):
                    alternate_website_scraped["number"] = scraped_from_google_website["number"]
                if scraped_from_google_website.get("scraped_from_link"):
                    alternate_website_scraped["scraped_from_link"] = scraped_from_google_website["scraped_from_link"]

    return idx, scraped, google_result, alternate_website_scraped


def _highlight_columns(buffer: io.BytesIO, all_cols: list, target_cols: set):
    """Apply a soft yellow fill + thin border to entire columns that are new."""
    yellow = PatternFill(fill_type="solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    buffer.seek(0)
    wb = load_workbook(buffer)
    ws = wb.active
    for col_idx, col_name in enumerate(all_cols, start=1):
        if col_name in target_cols:
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
                for cell in row:
                    cell.fill = yellow
                    cell.border = border
    buffer.truncate(0)
    buffer.seek(0)
    wb.save(buffer)

async def _scrape_website(website: str) -> dict:
    """Fallback: scrape the company website for email and number."""
    try:
        logger.info(f"scrape website for ")
        link_status, links, use_selenium = await fetch_links(website)
        if link_status == 0:
            return {}
        results = await extract_contents(links, use_selenium, website)
        if not results:
            return {}
        # Merge all found values, prefer first non-null
        email, number = None, None
        for r in results:
            scraped_from_link = r["link"]
            if not email and r.get("email"):
                email = r["email"]
            if not number and r.get("number"):
                number = r["number"]
        return {"email": email, "number": number,"scraped_from_link":scraped_from_link}
    except Exception as e:
        logger.error(f"Website scrape fallback failed for {website}: {e}")
        return {}

